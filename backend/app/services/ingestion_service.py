from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import pytesseract
from docx import Document as DocxDocument
from pdf2image import convert_from_bytes
from pypdf import PdfReader

from app.services.storage_service import safe_filename

# If native PDF text averages below this per page, fall back to OCR.
_OCR_FALLBACK_CHARS_PER_PAGE = 100


class IngestionError(RuntimeError):
    pass


class UnsupportedDocumentError(ValueError):
    pass


PDF_CONTENT_TYPES = {"application/pdf"}
DOCX_CONTENT_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
}


@dataclass(frozen=True)
class TextBlock:
    text: str
    page_number: int | None = None


@dataclass(frozen=True)
class TextChunk:
    text: str
    chunk_index: int
    page_number: int | None
    citation_label: str


def validate_supported_document(filename: str, content_type: str) -> None:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf" or content_type in PDF_CONTENT_TYPES:
        return
    if suffix == ".docx" or content_type in DOCX_CONTENT_TYPES:
        return
    raise UnsupportedDocumentError("Only PDF and DOCX uploads are supported")


def extract_text_blocks(file_bytes: bytes, filename: str, content_type: str) -> list[TextBlock]:
    validate_supported_document(filename, content_type)
    suffix = Path(filename).suffix.lower()

    if suffix == ".pdf" or content_type in PDF_CONTENT_TYPES:
        return extract_pdf_blocks(file_bytes)
    return extract_docx_blocks(file_bytes)


def extract_pdf_blocks(file_bytes: bytes) -> list[TextBlock]:
    blocks = _extract_pdf_native(file_bytes)
    page_count = max(len(blocks), 1)
    total_chars = sum(len(b.text) for b in blocks)

    if total_chars / page_count >= _OCR_FALLBACK_CHARS_PER_PAGE:
        print(f"[ingestion] native text extracted ({total_chars} chars, {len(blocks)} pages)")
        return blocks

    print(f"[ingestion] sparse native text ({total_chars} chars) — falling back to OCR")
    ocr_blocks = _extract_pdf_ocr(file_bytes)
    if not ocr_blocks:
        if blocks:
            return blocks
        raise IngestionError("No readable text found in the PDF (tried native extraction and OCR)")
    return ocr_blocks


def _extract_pdf_native(file_bytes: bytes) -> list[TextBlock]:
    try:
        reader = PdfReader(BytesIO(file_bytes))
        blocks = [
            TextBlock(text=(page.extract_text() or "").strip(), page_number=index)
            for index, page in enumerate(reader.pages, start=1)
        ]
    except Exception as exc:
        raise IngestionError(f"PDF text extraction failed: {exc}") from exc
    return [block for block in blocks if block.text]


def _extract_pdf_ocr(file_bytes: bytes) -> list[TextBlock]:
    # Fail fast if the required system binaries are absent rather than hanging.
    try:
        pytesseract.get_tesseract_version()
    except pytesseract.TesseractNotFoundError as exc:
        raise IngestionError("Tesseract is not installed — OCR unavailable") from exc

    print("[ingestion] OCR: converting PDF pages to images")
    try:
        images = convert_from_bytes(file_bytes, dpi=150)
    except Exception as exc:
        raise IngestionError(f"PDF to image conversion failed: {exc}") from exc

    print(f"[ingestion] OCR: running tesseract on {len(images)} pages")
    blocks: list[TextBlock] = []
    for page_num, image in enumerate(images, start=1):
        try:
            text = pytesseract.image_to_string(image, lang="eng").strip()
        except Exception as exc:
            raise IngestionError(f"OCR failed on page {page_num}: {exc}") from exc
        if text:
            blocks.append(TextBlock(text=text, page_number=page_num))
    print(f"[ingestion] OCR complete: {len(blocks)} pages with text")
    return blocks


def extract_docx_blocks(file_bytes: bytes) -> list[TextBlock]:
    try:
        document = DocxDocument(BytesIO(file_bytes))
        parts = [p.text.strip() for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if row_text:
                    parts.append(row_text)
    except Exception as exc:
        raise IngestionError(f"DOCX text extraction failed: {exc}") from exc

    text = "\n\n".join(p for p in parts if p)
    if not text:
        raise IngestionError("No readable text was found in the DOCX")
    return [TextBlock(text=text, page_number=None)]


def chunk_text_blocks(
    blocks: list[TextBlock],
    *,
    filename: str,
    chunk_size: int = 4000,
    overlap: int = 400,
) -> list[TextChunk]:
    if overlap >= chunk_size:
        raise ValueError("Chunk overlap must be smaller than chunk size")

    chunks: list[TextChunk] = []
    display_name = safe_filename(filename)

    for block in blocks:
        text = normalize_text(block.text)
        start = 0
        while start < len(text):
            end = min(start + chunk_size, len(text))
            chunk_text = text[start:end].strip()
            if chunk_text:
                chunk_index = len(chunks)
                chunks.append(
                    TextChunk(
                        text=chunk_text,
                        chunk_index=chunk_index,
                        page_number=block.page_number,
                        citation_label=make_citation_label(
                            display_name, block.page_number, chunk_index
                        ),
                    )
                )
            if end == len(text):
                break
            start = max(0, end - overlap)

    if not chunks:
        raise IngestionError("No chunks were produced from the extracted text")
    return chunks


def normalize_text(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.splitlines()]
    paragraphs = [line for line in lines if line]
    return "\n".join(paragraphs)


def make_citation_label(filename: str, page_number: int | None, chunk_index: int) -> str:
    if page_number is not None:
        return f"{filename} p. {page_number} ({chunk_index + 1})"
    return f"{filename} ({chunk_index + 1})"
